from __future__ import annotations

from pathlib import Path
import json
import re
import shutil

import openpyxl

XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")
SPELLS_JSON_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/data/rulesets/spells.json")


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def main() -> int:
    if not XLSX_PATH.exists():
        print(f"ERROR: workbook not found: {XLSX_PATH}")
        return 2
    if not SPELLS_JSON_PATH.exists():
        print(f"ERROR: spells.json not found: {SPELLS_JSON_PATH}")
        return 2

    with SPELLS_JSON_PATH.open("r", encoding="utf-8") as f:
        data = json.load(f)
    by_id = {t(s.get("id", "")).lower(): s for s in data.get("spells", []) if t(s.get("id", ""))}

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    if "ID" not in idx:
        print("ERROR: missing ID column")
        return 2

    # 1) Ensure Healing column exists immediately after Page.
    if "Healing" not in idx:
        if "Page" in idx:
            insert_at = idx["Page"] + 1
        else:
            insert_at = ws.max_column + 1
        ws.insert_cols(insert_at)
        ws.cell(1, insert_at).value = "Healing"

        # Fill default false for all populated rows
        for r in range(2, ws.max_row + 1):
            row_is_empty = True
            for c in range(1, ws.max_column + 1):
                if c == insert_at:
                    continue
                if t(ws.cell(r, c).value):
                    row_is_empty = False
                    break
            if not row_is_empty:
                ws.cell(r, insert_at).value = "false"

        print("Added Healing column and defaulted values to false.")
    else:
        hcol = idx["Healing"]
        filled = 0
        for r in range(2, ws.max_row + 1):
            if t(ws.cell(r, hcol).value) == "":
                ws.cell(r, hcol).value = "false"
                filled += 1
        if filled:
            print(f"Filled {filled} blank Healing cells with false.")

    # Rebuild header index after potential insert.
    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    # 2) Trim key text fields for clean import.
    trim_cols = [c for c in ["ID", "Category", "Level", "Name", "Schools", "Range", "Components", "Materials", "Page"] if c in idx]
    trim_updates = 0
    for r in range(2, ws.max_row + 1):
        for name in trim_cols:
            c = idx[name]
            old = ws.cell(r, c).value
            new = t(old)
            if old is not None and new != old:
                ws.cell(r, c).value = new
                trim_updates += 1
    if trim_updates:
        print(f"Trimmed whitespace in {trim_updates} cells.")

    # 3) Validate IDs and capture suspicious levels for manual review.
    id_col = idx["ID"]
    level_col = idx.get("Level")
    missing_ids = []
    suspicious_levels = []

    for r in range(2, ws.max_row + 1):
        rid = t(ws.cell(r, id_col).value)
        if not rid:
            continue

        if rid.lower() not in by_id:
            missing_ids.append((r, rid))

        if level_col:
            level = t(ws.cell(r, level_col).value)
            if level and not re.fullmatch(r"\d+|Quest", level):
                suspicious_levels.append((r, rid, level))

    # Backup then save in place
    backup = XLSX_PATH.with_suffix(".xlsx.bak")
    shutil.copy2(XLSX_PATH, backup)
    wb.save(XLSX_PATH)

    print(f"Saved normalized workbook: {XLSX_PATH}")
    print(f"Backup created: {backup}")
    print(f"Missing IDs vs spells.json: {len(missing_ids)}")
    print(f"Suspicious non-standard levels: {len(suspicious_levels)}")

    if missing_ids:
        print("Sample missing IDs:")
        for row, rid in missing_ids[:10]:
            print(f"  Row {row}: {rid}")

    if suspicious_levels:
        print("Sample suspicious levels:")
        for row, rid, level in suspicious_levels[:20]:
            print(f"  Row {row}: {rid} -> {level!r}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
