from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = ROOT / "Assets" / "Spells" / "Wizard Spells" / "Wizard_Spells_Tactical_Summaries_Standardized.xlsx"
OUTPUT_PATH = ROOT / "data" / "rulesets" / "spells.json"

FIELD_ALIASES = {
    "id": ["SpellID", "ID"],
    "category": ["Category"],
    "level": ["Level"],
    "name": ["Name"],
    "reversal": ["Reversal"],
    "schools": ["Schools"],
    "range": ["Range"],
    "components": ["Components"],
    "materials": ["Materials"],
    "cast_time": ["Cast Time"],
    "duration": ["Duration"],
    "area": ["Area"],
    "save": ["Save"],
    "frequency": ["Frequency"],
    "volume": ["Volume"],
    "page": ["Page"],
    "damage": ["Damage"],
    "damage_step": ["Damage Step"],
    "damage_scale_start_level": ["Start Level"],
    "damage_scale_every_levels": ["Every Levels"],
    "damage_max_at_level": ["Max At Level"],
    "damage_max": ["Max Damage"],
    "brief_description": ["Brief Description"],
    "description": ["Description"],
}

REQUIRED_FIELDS = {
    "id",
    "level",
    "name",
    "reversal",
    "schools",
    "range",
    "components",
    "materials",
    "cast_time",
    "duration",
    "area",
    "save",
    "frequency",
    "volume",
    "page",
    "damage",
    "damage_step",
    "damage_scale_start_level",
    "damage_scale_every_levels",
    "damage_max_at_level",
    "damage_max",
    "brief_description",
    "description",
}


def normalize(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_existing_spells() -> list[dict[str, str]]:
    if not OUTPUT_PATH.exists():
        return []

    data = json.loads(OUTPUT_PATH.read_text(encoding="utf-8"))
    spells = data.get("spells", [])
    if not isinstance(spells, list):
        return []
    return [spell for spell in spells if isinstance(spell, dict)]


def resolve_indexes(headers: list[str]) -> dict[str, int]:
    indexes = {header: idx for idx, header in enumerate(headers) if header}
    resolved: dict[str, int] = {}
    missing: list[str] = []

    for target, aliases in FIELD_ALIASES.items():
        idx = next((indexes[name] for name in aliases if name in indexes), None)
        if idx is None and target in REQUIRED_FIELDS:
            missing.append(f"{target} ({' | '.join(aliases)})")
            continue
        if idx is not None:
            resolved[target] = idx

    if missing:
        raise SystemExit("Missing expected columns: " + ", ".join(missing))

    return resolved


def extract_damage_config(text: str) -> dict[str, str]:
    lowered = (text or "").strip().lower()
    if not lowered:
        return {}

    # Ignore obvious healing-only wording to avoid false positives.
    if "hit points" in lowered and "damage" not in lowered and "suffer" not in lowered:
        return {}

    if not any(token in lowered for token in ("damage", "inflict", "suffer", "burn", "blast", "bolt", "fire", "cold", "acid", "electric")):
        return {}

    max_match = re.search(r"(?:up to|maximum(?: of)?|max(?:imum)?(?: of)?)\s*(\d+)d(\d+)", lowered)
    max_dice = f"{max_match.group(1)}d{max_match.group(2)}" if max_match else ""

    per_level_match = re.search(r"(\d+)d(\d+)\s*(?:points?\s+of\s+)?damage\s*(?:per|/)\s*level", lowered)
    if per_level_match:
        step = f"{per_level_match.group(1)}d{per_level_match.group(2)}"
        data = {
            "damage": step,
            "damage_step": step,
            "damage_scale_start_level": "1",
            "damage_scale_every_levels": "1",
        }
        if max_dice:
            data["damage_max"] = max_dice
            data["damage_max_at_level"] = ""
        return data

    per_n_levels_match = re.search(r"(\d+)d(\d+)\s*(?:points?\s+of\s+)?damage\s*(?:per|/)\s*(\d+)\s*levels?", lowered)
    if per_n_levels_match:
        step = f"{per_n_levels_match.group(1)}d{per_n_levels_match.group(2)}"
        data = {
            "damage": step,
            "damage_step": step,
            "damage_scale_start_level": per_n_levels_match.group(3),
            "damage_scale_every_levels": per_n_levels_match.group(3),
        }
        if max_dice:
            data["damage_max"] = max_dice
            data["damage_max_at_level"] = ""
        return data

    fixed_damage_match = re.search(r"(?:takes|take|inflicts?|causes?|suffers?)\s+(\d+)d(\d+)(?:\s*[+-]\s*\d+)?\s*(?:points?\s+of\s+)?damage", lowered)
    if fixed_damage_match:
        return {"damage": f"{fixed_damage_match.group(1)}d{fixed_damage_match.group(2)}"}

    generic_damage_match = re.search(r"\b(\d+)d(\d+)\b", lowered)
    if generic_damage_match and "damage" in lowered:
        return {"damage": f"{generic_damage_match.group(1)}d{generic_damage_match.group(2)}"}

    return {}


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK_PATH}")

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook["Spells"] if "Spells" in workbook.sheetnames else workbook[workbook.sheetnames[0]]

    rows = sheet.iter_rows(values_only=True)
    headers = [normalize(value) for value in next(rows)]
    indexes = resolve_indexes(headers)

    imported_arcane: list[dict[str, str]] = []
    derived_damage_count = 0
    for row in rows:
        spell_id = normalize(row[indexes["id"]])
        name = normalize(row[indexes["name"]])
        if not spell_id or not name:
            continue

        spell = {
            target: normalize(row[index])
            for target, index in indexes.items()
            if target != "category"
        }
        category = normalize(row[indexes["category"]]) if "category" in indexes else "arcane"
        spell["category"] = category.lower() if category else "arcane"
        spell["is_healing"] = False

        if not any(spell.get(key, "").strip() for key in (
            "damage",
            "damage_step",
            "damage_scale_start_level",
            "damage_scale_every_levels",
            "damage_max_at_level",
            "damage_max",
        )):
            extracted = extract_damage_config(
                " ".join([
                    spell.get("brief_description", ""),
                    spell.get("description", ""),
                ])
            )
            if extracted:
                spell.update(extracted)
                derived_damage_count += 1

        imported_arcane.append(spell)

    existing_non_arcane = [
        spell for spell in load_existing_spells()
        if normalize(spell.get("category", "")).lower() != "arcane"
    ]

    combined = imported_arcane + existing_non_arcane
    combined.sort(key=lambda spell: (
        normalize(spell.get("category", "")).lower(),
        normalize(spell.get("level", "")),
        normalize(spell.get("name", "")).lower(),
        normalize(spell.get("id", "")).lower(),
    ))

    backup_path = OUTPUT_PATH.with_suffix(".json.pre_wizard_reimport.bak")
    if OUTPUT_PATH.exists():
        shutil.copy2(OUTPUT_PATH, backup_path)

    OUTPUT_PATH.write_text(json.dumps({"spells": combined}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Imported {len(imported_arcane)} arcane spells to {OUTPUT_PATH}")
    print(f"Derived damage config for {derived_damage_count} arcane spells from text")
    print(f"Backup: {backup_path}")


if __name__ == "__main__":
    main()