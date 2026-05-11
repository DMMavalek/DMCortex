from __future__ import annotations

from pathlib import Path
import json
import shutil
import re

import openpyxl

XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")
JSON_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/data/rulesets/spells.json")


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def parse_bool(v: str) -> bool:
    s = v.strip().lower()
    return s in {"1", "true", "yes", "y"}


HEALING_NAME_PATTERN = re.compile(r"\b(cure|heal|healing|regenerate|restoration|restore)\b", re.IGNORECASE)


def has_damage_configuration(spell: dict[str, str]) -> bool:
    return any(
        bool((spell.get(key) or "").strip())
        for key in [
            "damage",
            "damage_step",
            "damage_max",
            "damage_max_at_level",
            "damage_scale_start_level",
            "damage_scale_every_levels",
        ]
    )


def infer_is_healing(spell: dict[str, str], explicit_healing: bool) -> bool:
    if explicit_healing:
        return True

    if not str(spell.get("category", "")).strip().lower() == "divine":
        return False

    if not has_damage_configuration(spell):
        return False

    name = str(spell.get("name", "")).strip()
    return bool(HEALING_NAME_PATTERN.search(name))


def main() -> int:
    if not XLSX_PATH.exists():
        print(f"ERROR: workbook not found: {XLSX_PATH}")
        return 2
    if not JSON_PATH.exists():
        print(f"ERROR: spells.json not found: {JSON_PATH}")
        return 2

    ws = openpyxl.load_workbook(XLSX_PATH).active
    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    required = [
        "ID", "Category", "Level", "Name", "Reversal", "Schools", "Range", "Components", "Materials",
        "Cast Time", "Duration", "Area", "Save", "Frequency", "Volume", "Page", "Healing",
        "Damage", "Damage Step", "Start Level", "Every Levels", "Max At Level", "Max Damage",
        "Brief Description", "Description",
    ]
    missing = [h for h in required if h not in idx]
    if missing:
        print("ERROR: missing required columns:")
        for h in missing:
            print(f"  - {h}")
        return 2

    with JSON_PATH.open("r", encoding="utf-8") as f:
        data = json.load(f)

    spells = data.get("spells", [])
    by_id = {t(s.get("id", "")).lower(): i for i, s in enumerate(spells) if t(s.get("id", ""))}

    updated = 0
    skipped = 0
    missing_ids = []
    inferred_healing_updates = 0

    for r in range(2, ws.max_row + 1):
        rid = t(ws.cell(r, idx["ID"]).value)
        if not rid:
            skipped += 1
            continue

        key = rid.lower()
        if key not in by_id:
            missing_ids.append((r, rid))
            skipped += 1
            continue

        i = by_id[key]
        explicit_healing = parse_bool(t(ws.cell(r, idx["Healing"]).value))

        row = {
            "id": rid,
            "category": t(ws.cell(r, idx["Category"]).value),
            "level": t(ws.cell(r, idx["Level"]).value),
            "name": t(ws.cell(r, idx["Name"]).value),
            "reversal": t(ws.cell(r, idx["Reversal"]).value),
            "schools": t(ws.cell(r, idx["Schools"]).value),
            "range": t(ws.cell(r, idx["Range"]).value),
            "components": t(ws.cell(r, idx["Components"]).value),
            "materials": t(ws.cell(r, idx["Materials"]).value),
            "cast_time": t(ws.cell(r, idx["Cast Time"]).value),
            "duration": t(ws.cell(r, idx["Duration"]).value),
            "area": t(ws.cell(r, idx["Area"]).value),
            "save": t(ws.cell(r, idx["Save"]).value),
            "frequency": t(ws.cell(r, idx["Frequency"]).value),
            "volume": t(ws.cell(r, idx["Volume"]).value),
            "page": t(ws.cell(r, idx["Page"]).value),
            "is_healing": explicit_healing,
            "damage": t(ws.cell(r, idx["Damage"]).value),
            "damage_step": t(ws.cell(r, idx["Damage Step"]).value),
            "damage_scale_start_level": t(ws.cell(r, idx["Start Level"]).value),
            "damage_scale_every_levels": t(ws.cell(r, idx["Every Levels"]).value),
            "damage_max_at_level": t(ws.cell(r, idx["Max At Level"]).value),
            "damage_max": t(ws.cell(r, idx["Max Damage"]).value),
            "brief_description": t(ws.cell(r, idx["Brief Description"]).value),
            "description": t(ws.cell(r, idx["Description"]).value),
        }

        if infer_is_healing(row, explicit_healing):
            if not row["is_healing"]:
                inferred_healing_updates += 1
            row["is_healing"] = True

        spells[i] = row
        updated += 1

    backup = JSON_PATH.with_suffix(".json.pre_priest_import.bak")
    shutil.copy2(JSON_PATH, backup)

    with JSON_PATH.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Updated spells: {updated}")
    print(f"Skipped rows: {skipped}")
    print(f"Inferred healing flags: {inferred_healing_updates}")
    print(f"Backup: {backup}")
    print(f"Missing IDs: {len(missing_ids)}")
    if missing_ids:
        for row, rid in missing_ids[:20]:
            print(f"  Row {row}: {rid}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
