from __future__ import annotations

from pathlib import Path
import json
import shutil

import openpyxl

JSON_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/data/rulesets/spells.json")
XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")

REPLACEMENTS = {
    "PRI001065": "Blood Bond",
    "PRI001069": "Beguiling",
}


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def update_json() -> int:
    data = json.load(JSON_PATH.open("r", encoding="utf-8"))
    spells = data.get("spells", [])
    by_id = {t(s.get("id", "")).upper(): s for s in spells}

    changed = 0
    for sid, new_name in REPLACEMENTS.items():
        s = by_id.get(sid)
        if s is None:
            continue
        if t(s.get("name", "")) != new_name:
            s["name"] = new_name
            changed += 1

    backup = JSON_PATH.with_suffix(".json.pre_final_two_names.bak")
    shutil.copy2(JSON_PATH, backup)
    json.dump(data, JSON_PATH.open("w", encoding="utf-8"), ensure_ascii=False, indent=2)
    return changed


def update_xlsx() -> int:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    id_col = idx["ID"]
    name_col = idx["Name"]

    changed = 0
    for r in range(2, ws.max_row + 1):
        sid = t(ws.cell(r, id_col).value).upper()
        if sid in REPLACEMENTS:
            new_name = REPLACEMENTS[sid]
            if t(ws.cell(r, name_col).value) != new_name:
                ws.cell(r, name_col).value = new_name
                changed += 1

    backup = XLSX_PATH.with_suffix(".xlsx.pre_final_two_names.bak")
    shutil.copy2(XLSX_PATH, backup)
    wb.save(XLSX_PATH)
    return changed


def main() -> int:
    json_changed = update_json()
    xlsx_changed = update_xlsx()
    print(f"JSON updated rows: {json_changed}")
    print(f"XLSX updated rows: {xlsx_changed}")
    print("Backups created:")
    print("- spells.json.pre_final_two_names.bak")
    print("- Priest_Spells_Complete.xlsx.pre_final_two_names.bak")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
