from __future__ import annotations

from pathlib import Path
import re

import openpyxl

TARGET = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")
SOURCE = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/AD_and_D_Priest_Spells_Final.xlsx")


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def norm_name(s: str) -> str:
    s = s.lower().strip()
    s = s.replace("’", "'")
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"[^a-z0-9' ]+", "", s)
    return s.strip()


def main() -> int:
    wt = openpyxl.load_workbook(TARGET)
    wst = wt.active
    ht = [t(wst.cell(1, c).value) for c in range(1, wst.max_column + 1)]
    it = {h: i + 1 for i, h in enumerate(ht) if h}

    ws = openpyxl.load_workbook(SOURCE).active
    hs = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    isrc = {h: i + 1 for i, h in enumerate(hs) if h}

    src_by_name: dict[str, str] = {}
    for r in range(2, ws.max_row + 1):
        name = t(ws.cell(r, isrc["Spell Name"]).value)
        lvl = t(ws.cell(r, isrc["Level"]).value)
        if not name:
            continue
        key = norm_name(name)
        if key and lvl:
            src_by_name[key] = lvl

    bad_rows = []
    for r in range(2, wst.max_row + 1):
        lvl = t(wst.cell(r, it["Level"]).value)
        if not lvl or re.fullmatch(r"\d+|Quest", lvl):
            continue
        rid = t(wst.cell(r, it["ID"]).value)
        name = t(wst.cell(r, it["Name"]).value)
        key = norm_name(name)
        mapped = src_by_name.get(key, "")
        bad_rows.append((r, rid, name, lvl, mapped))

    print(f"bad rows: {len(bad_rows)}")
    mapped_count = 0
    for row, rid, name, old, mapped in bad_rows:
        print(f"Row {row} | {rid} | {name} | old={old!r} | source={mapped!r}")
        if mapped:
            mapped_count += 1
    print(f"mapped from source: {mapped_count}/{len(bad_rows)}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
