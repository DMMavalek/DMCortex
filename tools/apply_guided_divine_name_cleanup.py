from __future__ import annotations

from pathlib import Path
import json
import shutil

import openpyxl

JSON_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/data/rulesets/spells.json")
XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")

# User-approved guided cleanup map.
REPLACEMENTS = {
    "PRI001063": "Animal Summoning III",
    "PRI001064": "Penetrate Disguise",
    "PRI001065": "UNKNOWN_Ela_Thieves",
    "PRI001066": "Cold Hand",
    "PRI001067": "Forgotten Melody",
    "PRI001068": "Watery Travel",
    "PRI001070": "Cure Rot",
    "PRI001071": "Animal Sight",
    "PRI001072": "Faith Magic Zone",
    "PRI001599": "Animate Statue",
    "PRI001747": "Call Stone Guardian",
}


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def update_json() -> tuple[int, list[str]]:
    data = json.load(JSON_PATH.open("r", encoding="utf-8"))
    spells = data.get("spells", [])
    by_id = {t(s.get("id", "")).upper(): s for s in spells}

    changed = 0
    missing = []
    for sid, new_name in REPLACEMENTS.items():
        s = by_id.get(sid)
        if s is None:
            missing.append(sid)
            continue
        old = t(s.get("name", ""))
        if old != new_name:
            s["name"] = new_name
            changed += 1

    backup = JSON_PATH.with_suffix(".json.pre_guided_name_cleanup.bak")
    shutil.copy2(JSON_PATH, backup)
    json.dump(data, JSON_PATH.open("w", encoding="utf-8"), ensure_ascii=False, indent=2)

    return changed, missing


def update_xlsx() -> tuple[int, list[str]]:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    if "ID" not in idx or "Name" not in idx:
        raise RuntimeError("Workbook missing ID or Name column")

    id_col = idx["ID"]
    name_col = idx["Name"]

    changed = 0
    seen = set()
    for r in range(2, ws.max_row + 1):
        sid = t(ws.cell(r, id_col).value).upper()
        if sid in REPLACEMENTS:
            seen.add(sid)
            new_name = REPLACEMENTS[sid]
            old = t(ws.cell(r, name_col).value)
            if old != new_name:
                ws.cell(r, name_col).value = new_name
                changed += 1

    missing = sorted(set(REPLACEMENTS.keys()) - seen)

    backup = XLSX_PATH.with_suffix(".xlsx.pre_guided_name_cleanup.bak")
    shutil.copy2(XLSX_PATH, backup)
    wb.save(XLSX_PATH)

    return changed, missing


def main() -> int:
    json_changed, json_missing = update_json()
    xlsx_changed, xlsx_missing = update_xlsx()

    print(f"JSON names updated: {json_changed}")
    print(f"XLSX names updated: {xlsx_changed}")
    if json_missing:
        print("Missing IDs in JSON:", ", ".join(json_missing))
    if xlsx_missing:
        print("Missing IDs in XLSX:", ", ".join(xlsx_missing))

    print("Backups created:")
    print("- spells.json.pre_guided_name_cleanup.bak")
    print("- Priest_Spells_Complete.xlsx.pre_guided_name_cleanup.bak")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
