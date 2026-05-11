from __future__ import annotations

from pathlib import Path
import re
import shutil

import openpyxl

XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def normalize_level(raw: str) -> str:
    s = raw.strip()
    if not s:
        return s
    if re.fullmatch(r"\d+|Quest", s):
        return s
    if s == "|":
        return "1"

    m = re.search(r"([1-7])", s)
    if m:
        return m.group(1)

    return s


def main() -> int:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    if "Level" not in idx or "ID" not in idx:
        print("ERROR: missing required Level or ID column")
        return 2

    changed = []
    col = idx["Level"]
    id_col = idx["ID"]

    for r in range(2, ws.max_row + 1):
        old = t(ws.cell(r, col).value)
        if not old:
            continue
        new = normalize_level(old)
        if new != old:
            rid = t(ws.cell(r, id_col).value)
            ws.cell(r, col).value = new
            changed.append((r, rid, old, new))

    backup = XLSX_PATH.with_suffix(".xlsx.levels.bak")
    shutil.copy2(XLSX_PATH, backup)
    wb.save(XLSX_PATH)

    print(f"Saved normalized levels: {XLSX_PATH}")
    print(f"Backup created: {backup}")
    print(f"Rows changed: {len(changed)}")
    for row, rid, old, new in changed[:50]:
        print(f"  Row {row} | {rid}: {old!r} -> {new!r}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
