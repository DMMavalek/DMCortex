from __future__ import annotations

from pathlib import Path
import re

import openpyxl

XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")


def t(v: object) -> str:
    return "" if v is None else str(v).strip()


def main() -> int:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    headers = [t(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    required = ["ID", "Level", "Name", "Brief Description", "Description"]
    for col in required:
        if col not in idx:
            print(f"Missing column: {col}")
            return 2

    bad = []
    for r in range(2, ws.max_row + 1):
        level = t(ws.cell(r, idx["Level"]).value)
        if not level:
            continue
        if re.fullmatch(r"\d+|Quest", level):
            continue

        rid = t(ws.cell(r, idx["ID"]).value)
        name = t(ws.cell(r, idx["Name"]).value)
        brief = t(ws.cell(r, idx["Brief Description"]).value)
        desc = t(ws.cell(r, idx["Description"]).value)

        source_text = f"{brief} {desc}".strip()
        ord_match = re.search(r"\b([1-7])(?:st|nd|rd|th)\b", source_text, re.IGNORECASE)
        digit_match = re.search(r"\b([1-7])\b", source_text)

        hint = ""
        if ord_match:
            hint = ord_match.group(1)
        elif digit_match:
            hint = digit_match.group(1)

        bad.append((r, rid, name, level, hint, (brief or desc)[:160]))

    print(f"Suspicious levels: {len(bad)}")
    for row, rid, name, level, hint, snippet in bad:
        print("---")
        print(f"Row {row} | {rid} | {name}")
        print(f"Level raw: {level!r} | Hint: {hint!r}")
        print(f"Snippet: {snippet}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
