from __future__ import annotations

from collections import Counter
from pathlib import Path
import json
import re

import openpyxl

XLSX_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/Assets/Spells/Priest Spells/Priest_Spells_Complete.xlsx")
SPELLS_JSON_PATH = Path(r"C:/Users/kelava/Documents/Projects/Dungeon Master Cortex/data/rulesets/spells.json")

EXPECTED_HEADERS = [
    "ID", "Category", "Level", "Name", "Reversal", "Schools",
    "Range", "Components", "Materials", "Cast Time", "Duration", "Area",
    "Save", "Frequency", "Volume", "Page", "Healing", "Damage", "Damage Step",
    "Start Level", "Every Levels", "Max At Level", "Max Damage",
    "Brief Description", "Description",
]

BOOL_TRUE = {"true", "1", "yes", "y"}
BOOL_FALSE = {"false", "0", "no", "n", ""}
HEALING_NAME_PATTERN = re.compile(r"\b(cure|heal|healing|regenerate|restoration|restore)\b", re.IGNORECASE)


def normalize_text(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def has_damage_configuration(row: dict[str, str]) -> bool:
    return any(
        bool(row.get(key, "").strip())
        for key in ["Damage", "Damage Step", "Max Damage", "Start Level", "Every Levels", "Max At Level"]
    )


def validate() -> int:
    if not XLSX_PATH.exists():
        print(f"ERROR: Workbook not found: {XLSX_PATH}")
        return 2

    if not SPELLS_JSON_PATH.exists():
        print(f"ERROR: spells.json not found: {SPELLS_JSON_PATH}")
        return 2

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active

    headers = [normalize_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    header_set = set(headers)

    print(f"Workbook: {XLSX_PATH.name}")
    print(f"Sheet: {ws.title}")
    print(f"Rows (incl header): {ws.max_row}")
    print(f"Columns: {ws.max_column}")

    missing_headers = [h for h in EXPECTED_HEADERS if h not in header_set]
    extra_headers = [h for h in headers if h and h not in EXPECTED_HEADERS]

    print("\nHeader check:")
    print(f"- Missing expected headers: {len(missing_headers)}")
    if missing_headers:
        for h in missing_headers:
            print(f"  * {h}")
    print(f"- Unexpected extra headers: {len(extra_headers)}")
    if extra_headers:
        for h in extra_headers:
            print(f"  * {h}")

    col_idx = {h: i + 1 for i, h in enumerate(headers) if h}

    with SPELLS_JSON_PATH.open("r", encoding="utf-8-sig") as f:
        data = json.load(f)
    spells = data.get("spells", [])
    spell_ids = {str(s.get("id", "")).strip().lower() for s in spells if str(s.get("id", "")).strip()}

    issues = Counter()
    duplicate_ids: dict[str, list[int]] = {}
    id_rows: dict[str, list[int]] = {}

    last_row = ws.max_row
    for r in range(2, last_row + 1):
        row_vals = [normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if all(v == "" for v in row_vals):
            issues["blank_rows"] += 1
            continue

        rid = normalize_text(ws.cell(r, col_idx.get("ID", 0)).value) if "ID" in col_idx else ""
        category = normalize_text(ws.cell(r, col_idx.get("Category", 0)).value) if "Category" in col_idx else ""
        name = normalize_text(ws.cell(r, col_idx.get("Name", 0)).value) if "Name" in col_idx else ""
        healing = normalize_text(ws.cell(r, col_idx.get("Healing", 0)).value) if "Healing" in col_idx else ""

        row_map = {
            "Damage": normalize_text(ws.cell(r, col_idx.get("Damage", 0)).value) if "Damage" in col_idx else "",
            "Damage Step": normalize_text(ws.cell(r, col_idx.get("Damage Step", 0)).value) if "Damage Step" in col_idx else "",
            "Start Level": normalize_text(ws.cell(r, col_idx.get("Start Level", 0)).value) if "Start Level" in col_idx else "",
            "Every Levels": normalize_text(ws.cell(r, col_idx.get("Every Levels", 0)).value) if "Every Levels" in col_idx else "",
            "Max At Level": normalize_text(ws.cell(r, col_idx.get("Max At Level", 0)).value) if "Max At Level" in col_idx else "",
            "Max Damage": normalize_text(ws.cell(r, col_idx.get("Max Damage", 0)).value) if "Max Damage" in col_idx else "",
        }

        if not rid:
            issues["missing_id"] += 1
        else:
            key = rid.lower()
            id_rows.setdefault(key, []).append(r)
            if key not in spell_ids:
                issues["id_not_in_spells_json"] += 1

        if not name:
            issues["missing_name"] += 1

        if category.lower() != "divine":
            issues["non_divine_category"] += 1

        if healing.lower() not in BOOL_TRUE and healing.lower() not in BOOL_FALSE:
            issues["invalid_healing_value"] += 1

        # Warn-only heuristic: likely healing spell names with numeric effect config but Healing=false/blank.
        if (
            category.lower() == "divine"
            and HEALING_NAME_PATTERN.search(name)
            and has_damage_configuration(row_map)
            and healing.lower() in BOOL_FALSE
        ):
            issues["likely_healing_marked_false"] += 1

        # Light validation on level and integer scaling fields
        level = normalize_text(ws.cell(r, col_idx.get("Level", 0)).value) if "Level" in col_idx else ""
        if level and not re.fullmatch(r"\d+", level):
            issues["non_numeric_level"] += 1

        for fld in ["Start Level", "Every Levels", "Max At Level"]:
            if fld in col_idx:
                v = normalize_text(ws.cell(r, col_idx[fld]).value)
                if v and not re.fullmatch(r"\d+", v):
                    issues[f"invalid_int_{fld.lower().replace(' ', '_')}"] += 1

        # Optional damage format check (allows empty)
        for fld in ["Damage", "Damage Step", "Max Damage"]:
            if fld in col_idx:
                v = normalize_text(ws.cell(r, col_idx[fld]).value)
                if v and not re.fullmatch(r"[0-9dD+\- ]+", v):
                    issues[f"odd_damage_format_{fld.lower().replace(' ', '_')}"] += 1

    for sid, rows in id_rows.items():
        if len(rows) > 1:
            duplicate_ids[sid] = rows

    print("\nRow validation summary:")
    tracked = [
        "blank_rows", "missing_id", "missing_name", "non_divine_category",
        "id_not_in_spells_json", "invalid_healing_value", "non_numeric_level",
        "likely_healing_marked_false",
        "invalid_int_start_level", "invalid_int_every_levels", "invalid_int_max_at_level",
        "odd_damage_format_damage", "odd_damage_format_damage_step", "odd_damage_format_max_damage",
    ]
    for k in tracked:
        print(f"- {k}: {issues.get(k, 0)}")

    print(f"- duplicate_id_groups: {len(duplicate_ids)}")
    if duplicate_ids:
        print("  Sample duplicates:")
        for i, (sid, rows) in enumerate(duplicate_ids.items()):
            print(f"  * {sid}: rows {rows[:6]}")
            if i >= 9:
                break

    blockers = (
        len(missing_headers)
        + issues.get("missing_id", 0)
        + len(duplicate_ids)
    )

    print("\nReadiness:")
    if blockers == 0:
        print("- READY for reimport (no structural blockers found).")
        return 0

    print("- NOT READY for reimport (structural blockers found).")
    return 1


if __name__ == "__main__":
    raise SystemExit(validate())
